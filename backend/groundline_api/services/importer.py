"""CSV / XLSX / JSON / JSONL / YAML import: parsing, mapping, validation (§4, §8).

Two-step flow used by `api/v1/imports.py`:
  - `preview()` -- a single pass over the file that returns a sample of the
    first rows, the total row count, and (if a mapping is supplied) the full
    per-row validation report. Persists nothing.
  - `commit()` -- a single pass that applies the mapping plus caller-supplied
    skip/fix overrides and returns, per source row, whether it imported,
    was explicitly skipped, or errored. The route creates a `DatasetRow` for
    each row `commit()` returns as valid.

CSV/XLSX are streamed (the CSV reader / openpyxl `iter_rows` generator,
consumed once). JSON/JSONL/YAML are parsed whole by the shared
`groundline_schema.records` (GL-3.5-5), and each record is flattened to dotted
source columns, with the dataset's `json` column keys kept as leaves.

Row index convention: 1-based, counting only rows of data (the header row of
a CSV/XLSX is not numbered; for JSON/JSONL/YAML it's the record's position).

Coercion from spreadsheet strings, applied per mapped column before
validation (GL-2-8):
  - number:        "4" / "4.5" -> int / float; anything else that doesn't
                    parse is left as the raw string, so validation reports
                    "expected a number".
  - boolean:        true/false/1/0/yes/no (case-insensitive, trimmed) -> bool;
                    anything else is left as the raw string.
  - multi_select:   split on comma or semicolon, each part trimmed, empty
                    parts dropped.
  - select:         trimmed string (still option-checked downstream).
  - text/long_text: the raw string, unmodified.
  - json:           the raw string, as a JSON string.
An empty cell is treated as an absent value (the column's `required` check
applies to it; it is never coerced).

Dates (GL-3.5-9), every format: a `YYYY-MM-DD` string is kept; a string or
XLSX cell with a time component is refused ("has a time component") unless
it's an XLSX datetime at exactly midnight; XLSX time/duration cells and
numbers are refused (a serial or epoch number is never guessed); any other
string is left for validation to report. XLSX date cells stay native values
(read with the workbook's own 1900/1904 epoch) instead of their str().

Typed values from JSON/YAML (GL-3.5-5): strings -- and plain YAML scalars,
by their raw text -- take the string rules above, except that a plain YAML
scalar going to a `json` column is typed the way JSON would type it. Numbers
and booleans are kept where the column type allows (number -> text is its
string; 0/1 -> boolean; booleans -> "true"/"false" text); arrays fit only
multi_select or json, objects only json. null means absent.

If the dataset has a key column, a key repeated within the file or already
used in the dataset is a row error; import never updates rows (that's sync).
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import zipfile
from collections.abc import Iterator
from dataclasses import dataclass, field
from typing import Any

import openpyxl
from groundline_schema import ColumnType
from groundline_schema.paths import PathError, flatten
from groundline_schema.records import (
    STRUCTURED_SUFFIXES,
    PlainScalar,
    RecordProblem,
    RecordsError,
    decode_text,
    parse_records,
    to_json_value,
)
from groundline_schema.temporal import is_date
from openpyxl.utils.exceptions import InvalidFileException

from groundline_api.config import settings
from groundline_api.models.dataset import DatasetColumn
from groundline_api.services.validation import key_column, row_errors

SAMPLE_SIZE = 20


class ImportRequestError(Exception):
    """A caller-facing (422) error in the uploaded file or the supplied mapping."""

    def __init__(self, message: str, candidates: list[str] | None = None) -> None:
        super().__init__(message)
        self.candidates = candidates  # records_key choices, when one is needed


class _CoercionError(Exception):
    """A value that can't go into its mapped column at all (a row error)."""


@dataclass
class ParsedSource:
    header: list[str]
    # Per data row: {source column: value}, or a RecordProblem for a record
    # that couldn't be read. CSV/XLSX values are raw cell strings.
    rows: Iterator[dict[str, Any] | RecordProblem]
    structured: bool = False
    records_key: str | None = None
    ignored_keys: list[str] = field(default_factory=list)


def parse_upload(
    filename: str,
    content: bytes,
    columns: list[DatasetColumn] | None = None,
    records_key: str | None = None,
) -> ParsedSource:
    """Parse an upload into source columns and a row iterator."""
    lower = filename.lower()
    if lower.endswith(".csv"):
        return _parse_csv(content)
    if lower.endswith(".xlsx"):
        return _parse_xlsx(content)
    if lower.endswith(STRUCTURED_SUFFIXES):
        return _parse_structured(filename, content, columns or [], records_key)
    raise ImportRequestError(
        f"unsupported file type: {filename!r} "
        "(expected .csv, .xlsx, .json, .jsonl, .ndjson, .yaml or .yml)"
    )


def _parse_csv(content: bytes) -> ParsedSource:
    try:
        text = decode_text(content, csv=True)
    except RecordsError as exc:
        raise ImportRequestError(str(exc)) from None
    reader = csv.reader(io.StringIO(text))
    try:
        header = next(reader)
    except StopIteration:
        header = []

    def rows() -> Iterator[dict[str, Any]]:
        for raw_row in reader:
            yield _cells(header, raw_row)

    return ParsedSource(header=header, rows=rows())


def _parse_xlsx(content: bytes) -> ParsedSource:
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        row_iter = workbook.worksheets[0].iter_rows(values_only=True)
        first = next(row_iter, None)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError, IndexError):
        raise ImportRequestError(
            "could not read the workbook — is it a valid, unencrypted .xlsx file?"
        ) from None
    header = [] if first is None else ["" if cell is None else str(cell) for cell in first]

    def rows() -> Iterator[dict[str, Any]]:
        for raw_row in row_iter:
            yield _cells(header, [_xlsx_cell(cell) for cell in raw_row])

    return ParsedSource(header=header, rows=rows())


def _xlsx_cell(cell: Any) -> Any:
    # Date/time cells stay native so date columns can tell a date from a
    # datetime; every other column still sees their str(), as before.
    if cell is None:
        return ""
    if isinstance(cell, (dt.date, dt.time, dt.timedelta)):
        return cell
    return str(cell)


def _cells(header: list[str], raw_row: list[Any]) -> dict[str, Any]:
    return {header[i]: (raw_row[i] if i < len(raw_row) else "") for i in range(len(header))}


def _parse_structured(
    filename: str, content: bytes, columns: list[DatasetColumn], records_key: str | None
) -> ParsedSource:
    try:
        parsed = parse_records(filename, content, records_key, settings.import_max_nodes)
    except RecordsError as exc:
        raise ImportRequestError(str(exc), exc.candidates) from None

    leaves = {col.key for col in columns if col.type == ColumnType.JSON}
    header: list[str] = []
    seen: set[str] = set()
    rows: list[dict[str, Any] | RecordProblem] = []
    for record in parsed.records:
        if isinstance(record, RecordProblem):
            rows.append(record)
            continue
        try:
            flat = flatten(record, leaves)
        except PathError as exc:
            rows.append(RecordProblem(str(exc)))
            continue
        for path in flat:
            if path not in seen:
                seen.add(path)
                header.append(path)
        rows.append(flat)

    return ParsedSource(
        header=header,
        rows=iter(rows),
        structured=True,
        records_key=parsed.records_key,
        ignored_keys=parsed.ignored_keys,
    )


def _iterate_rows(source: ParsedSource) -> Iterator[tuple[int, dict[str, Any] | RecordProblem]]:
    """Yield (1-based row index, row values or a RecordProblem) for each data row."""
    yield from enumerate(source.rows, start=1)


def suggest_mapping(header: list[str], columns: list[DatasetColumn]) -> dict[str, str]:
    """Case-insensitive header -> schema key matches; first match wins per key."""
    keys_by_lower = {col.key.lower(): col.key for col in columns if not col.archived}
    mapping: dict[str, str] = {}
    used_keys: set[str] = set()
    for source_col in header:
        key = keys_by_lower.get(source_col.strip().lower())
        if key and key not in used_keys:
            mapping[source_col] = key
            used_keys.add(key)
    return mapping


def validate_mapping(mapping: dict[str, str], columns: list[DatasetColumn]) -> None:
    """Raise ImportRequestError for a mapping that references an unknown or duplicate key."""
    valid_keys = {col.key for col in columns}
    seen_keys: dict[str, str] = {}
    for source_col, schema_key in mapping.items():
        if schema_key not in valid_keys:
            raise ImportRequestError(f"mapping: '{schema_key}' is not a schema column")
        if schema_key in seen_keys:
            raise ImportRequestError(
                f"mapping: '{seen_keys[schema_key]}' and '{source_col}' "
                f"both map to '{schema_key}'"
            )
        seen_keys[schema_key] = source_col


def _coerce(raw: str, col_type: ColumnType) -> Any:
    if col_type == ColumnType.SELECT:
        return raw.strip()
    if col_type == ColumnType.MULTI_SELECT:
        return [part.strip() for part in re.split(r"[,;]", raw) if part.strip()]
    if col_type == ColumnType.NUMBER:
        try:
            return int(raw)
        except ValueError:
            try:
                return float(raw)
            except ValueError:
                return raw
    if col_type == ColumnType.BOOLEAN:
        lowered = raw.strip().lower()
        if lowered in ("true", "1", "yes"):
            return True
        if lowered in ("false", "0", "no"):
            return False
        return raw
    return raw  # text / long_text / json: as-is


_HAS_TIME = re.compile(r"\d{4}-\d{2}-\d{2}[T ]\d")


def _coerce_date(value: Any) -> Any:
    """Coerce one value for a date column (see the module docstring)."""
    if isinstance(value, dt.datetime):
        if value.time() != dt.time(0):
            raise _CoercionError(f"{value.isoformat(sep=' ')} has a time component")
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, (dt.time, dt.timedelta)):
        raise _CoercionError("a time is not a date")
    if isinstance(value, (bool, int, float)):
        raise _CoercionError(f"{value!r} is a number, not a date (write YYYY-MM-DD)")
    if isinstance(value, str):
        text = value.strip()
        if is_date(text):
            return text
        if _HAS_TIME.match(text):
            raise _CoercionError(f"{text!r} has a time component")
        return str(value)  # validation names the expected YYYY-MM-DD form
    raise _CoercionError("map this to a json column")


def _coerce_typed(value: Any, col_type: ColumnType) -> Any:
    """Coerce one JSON/YAML value for its column (see the module docstring)."""
    if col_type == ColumnType.DATE:
        return _coerce_date(value)
    if col_type == ColumnType.JSON:
        try:
            return to_json_value(value)
        except ValueError as exc:
            raise _CoercionError(str(exc)) from None
    if isinstance(value, str):  # including plain YAML scalars, by their raw text
        return _coerce(str(value), col_type)
    if isinstance(value, bool):
        if col_type in (ColumnType.TEXT, ColumnType.LONG_TEXT, ColumnType.SELECT):
            return "true" if value else "false"
        return value
    if isinstance(value, (int, float)):
        if col_type in (ColumnType.TEXT, ColumnType.LONG_TEXT, ColumnType.SELECT):
            return str(value)
        if col_type == ColumnType.BOOLEAN and value in (0, 1):
            return bool(value)
        return value
    if isinstance(value, list) and col_type == ColumnType.MULTI_SELECT:
        return [str(item) if isinstance(item, PlainScalar) else item for item in value]
    raise _CoercionError("map this to a json column")


def map_row(
    raw_values: dict[str, Any],
    mapping: dict[str, str],
    columns_by_key: dict[str, DatasetColumn],
    structured: bool,
) -> tuple[dict[str, Any], list[tuple[str, str]]]:
    """Apply the mapping + type coercion to one source row.

    Returns (data, coercion failures as (column_key, reason)).
    """
    data: dict[str, Any] = {}
    failures: list[tuple[str, str]] = []
    for source_col, schema_key in mapping.items():
        raw = raw_values.get(source_col)
        col_type = columns_by_key[schema_key].type
        if not structured:
            if raw is None or (isinstance(raw, str) and raw.strip() == ""):
                continue  # empty cell = absent
            try:
                if col_type == ColumnType.DATE:
                    data[schema_key] = _coerce_date(raw)
                else:  # an XLSX date cell outside a date column: its str(), as before
                    data[schema_key] = _coerce(raw if isinstance(raw, str) else str(raw), col_type)
            except _CoercionError as exc:
                failures.append((schema_key, str(exc)))
            continue
        if raw is None:
            continue  # null / missing = absent
        try:
            data[schema_key] = _coerce_typed(raw, col_type)
        except _CoercionError as exc:
            failures.append((schema_key, str(exc)))
    return data, failures


def _sample_values(values: dict[str, Any] | RecordProblem) -> dict[str, Any]:
    """A row's source values for the preview sample, typed as JSON would show them."""
    if isinstance(values, RecordProblem):
        return {}
    shown: dict[str, Any] = {}
    for column, value in values.items():
        try:
            shown[column] = to_json_value(value)
        except ValueError:
            shown[column] = str(value)
    return shown


@dataclass
class _KeyCheck:
    """Key column rules for one pass: unique within the file, new to the dataset."""

    column: DatasetColumn | None
    existing: set[str]
    first_row: dict[str, int] = field(default_factory=dict)

    def failures(self, index: int, data: dict[str, Any]) -> list[tuple[str, str]]:
        if self.column is None:
            return []
        value = data.get(self.column.key)
        if not isinstance(value, str) or value.strip() == "":
            return []  # row_errors already reports a missing/invalid key
        key = self.column.key
        if value in self.first_row:
            return [(key, f"duplicate key '{value}' in file (also row {self.first_row[value]})")]
        self.first_row[value] = index
        if value in self.existing:
            return [(key, f"key '{value}' already exists — use rows push to update")]
        return []


def _validate_row(
    index: int,
    values: dict[str, Any] | RecordProblem,
    mapping: dict[str, str],
    columns: list[DatasetColumn],
    columns_by_key: dict[str, DatasetColumn],
    structured: bool,
    fixes: dict[str, Any],
    keys: _KeyCheck,
) -> tuple[dict[str, Any], list[tuple[str, str]]]:
    if isinstance(values, RecordProblem):
        return {}, [("", values.reason)]
    data, failures = map_row(values, mapping, columns_by_key, structured)
    for key, value in fixes.items():
        # A fix replaces the source value, so its coercion failure no longer applies;
        # date fixes take the same date rules as the source.
        failures = [(k, reason) for k, reason in failures if k != key]
        column = columns_by_key.get(key)
        try:
            data[key] = _coerce_date(value) if column and column.type == ColumnType.DATE and value is not None else value
        except _CoercionError as exc:
            failures.append((key, str(exc)))
    failed = {column for column, _ in failures}
    failures += [(k, reason) for k, reason in row_errors(data, columns) if k not in failed]
    failures += keys.failures(index, data)
    return data, failures


def preview(
    source: ParsedSource,
    columns: list[DatasetColumn],
    mapping: dict[str, str] | None,
    existing_keys: set[str] | None = None,
) -> tuple[list[dict[str, Any]], int, dict[str, Any] | None]:
    """Single pass: sample rows, total count, and an optional validation report.

    Returns (sample_rows, total_rows, validation) where sample_rows is a list
    of `{"row": index, "values": {...}}` (up to SAMPLE_SIZE), and validation
    is `None` if no mapping was supplied, else `{"valid": n, "errors": [...]}`.
    """
    columns_by_key = {col.key: col for col in columns}
    keys = _KeyCheck(key_column(columns), existing_keys or set())
    sample: list[dict[str, Any]] = []
    total = 0
    valid = 0
    errors: list[dict[str, Any]] = []

    for index, values in _iterate_rows(source):
        total += 1
        if len(sample) < SAMPLE_SIZE:
            sample.append({"row": index, "values": _sample_values(values)})
        if mapping is not None:
            _, failures = _validate_row(
                index, values, mapping, columns, columns_by_key, source.structured, {}, keys
            )
            if failures:
                errors.extend(
                    {"row": index, "column": key, "reason": reason} for key, reason in failures
                )
            else:
                valid += 1

    validation = None if mapping is None else {"valid": valid, "errors": errors}
    return sample, total, validation


def commit(
    source: ParsedSource,
    columns: list[DatasetColumn],
    mapping: dict[str, str],
    skip: set[int],
    fixes: dict[int, dict[str, Any]],
    existing_keys: set[str] | None = None,
) -> tuple[list[dict[str, Any]], int, int, list[dict[str, Any]]]:
    """Single pass: apply mapping + skip/fixes, validate, and report the outcome.

    Returns (row_data_to_create, imported_count, skipped_count, errors). Every
    source row appears in exactly one of: the created rows, the skipped
    count, or the errors list (never silently dropped).
    """
    columns_by_key = {col.key: col for col in columns}
    keys = _KeyCheck(key_column(columns), existing_keys or set())
    created: list[dict[str, Any]] = []
    imported = 0
    skipped = 0
    errors: list[dict[str, Any]] = []

    for index, values in _iterate_rows(source):
        if index in skip:
            skipped += 1
            continue
        data, failures = _validate_row(
            index,
            values,
            mapping,
            columns,
            columns_by_key,
            source.structured,
            fixes.get(index, {}),
            keys,
        )
        if failures:
            errors.extend(
                {"row": index, "column": key, "reason": reason} for key, reason in failures
            )
            continue
        created.append(data)
        imported += 1

    return created, imported, skipped, errors
